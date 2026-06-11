import io

from openpyxl import load_workbook

from quanquant.stats.export import to_csv_bytes, to_xlsx_bytes
from quanquant.stats.metrics import compute_stats


def test_csv_has_header_and_rows(sample_trades):
    text = to_csv_bytes(sample_trades).decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert "商品" in lines[0]
    assert len(lines) == len(sample_trades) + 1  # header + one row per trade


def test_xlsx_has_two_sheets(sample_trades):
    closed = [t for t in sample_trades if t.exit_time is not None]
    data = to_xlsx_bytes(sample_trades, compute_stats(closed))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["交易明細", "績效統計"]
    assert wb["交易明細"].max_row == len(sample_trades) + 1
